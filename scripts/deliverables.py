"""Generate the two-page architecture PDF and five-slide technical presentation."""

import json
import math
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "deliverables"
OUT.mkdir(exist_ok=True)


def measured(file, field, fallback="Pending qualification"):
    path = ROOT / "runtime" / file
    value = json.loads(path.read_text()) if path.exists() else {}
    for part in field.split("."):
        value = value.get(part, fallback) if isinstance(value, dict) else fallback
    return value


def architecture():
    pdf = canvas.Canvas(str(OUT / "architecture.pdf"), pagesize=A4)
    width, height = A4

    def title(number, heading):
        pdf.setFillColor(HexColor("#17253d"))
        pdf.rect(0, height - 105, width, 105, fill=1, stroke=0)
        pdf.setFillColor(HexColor("#ffffff"))
        pdf.setFont("Helvetica-Bold", 22)
        pdf.drawString(40, height - 50, heading)
        pdf.setFont("Helvetica", 10)
        pdf.drawString(
            40, height - 76, "PROOFLANE  |  Configuration Evidence and Review Workspace"
        )
        pdf.setFillColor(HexColor("#52627a"))
        pdf.setFont("Helvetica", 9)
        pdf.drawString(40, 25, f"Local Docker implementation | Page {number} / 2 | 28 September 2026")

    def block(y, heading, lines):
        pdf.setFillColor(HexColor("#17253d"))
        pdf.setFont("Helvetica-Bold", 13)
        pdf.drawString(40, y, heading)
        pdf.setFont("Helvetica", 10)
        pdf.setFillColor(HexColor("#43536b"))
        for index, line in enumerate(lines):
            pdf.drawString(40, y - 22 - index * 16, line)

    title(1, "Evidence before verdicts")
    block(
        705,
        "Purpose and adaptation mechanism",
        [
            "Upload configurations, normalize supported facts, evaluate reviewed controls and export evidence.",
            "Unfamiliar syntax enters a bounded Gemini investigation: retrieve references, propose a mapping,",
            "test positive/negative cases, request clarification, and require reviewer activation.",
            "A new audit uses the approved mapping without code deployment. Prior versions stay pinned.",
        ],
    )
    block(
        588,
        "Runtime architecture",
        [
            "Browser -> NGINX gateway -> stateless FastAPI replicas -> PostgreSQL / Valkey / encrypted S3.",
            "Keycloak provides OIDC with PKCE; server-side sessions carry trusted organization membership.",
            "Worker replicas claim durable jobs with leases and fencing. A separate stage creates device PDFs.",
            "Gemini requests pass through a proxy restricted to Google inference endpoints.",
            "Maintenance expires artifacts, completes deletion and reconciles orphaned objects.",
            "Optional Prometheus/Grafana and OTLP profiles expose operational evidence.",
        ],
    )
    block(
        438,
        "Data and service ownership",
        [
            "PostgreSQL: tenant records, immutable input/policy/mapping references, sessions, jobs and events.",
            "SeaweedFS S3: AES-GCM encrypted source snapshots and reports, with SHA-256 integrity checks.",
            "Valkey: distributed token buckets, provider concurrency permits and token reservations.",
            "One shared Python domain package keeps worker logic consistent across independent processes.",
        ],
    )
    block(
        318,
        "Source and interpretation boundaries",
        [
            "20 team-authored technical checks; initial IOS/IOS-XE, Junos set and FortiOS subsets.",
            "Text/JSON/XML mappings have fixed selectors and types; generated code is never executed.",
            "Unknown, conflicting or unsupported effective settings never silently become a pass.",
            "Shipped configuration samples are synthetic; reference documents require authorized use.",
            "Framework selections are candidate crosswalks, not full certified CIS/NIST/STIG/ISO packs.",
        ],
    )
    pdf.setFillColor(HexColor("#17253d"))
    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawString(40, 191, "Local service connections")

    def box(x, y, w, label):
        pdf.setStrokeColor(HexColor("#c5d1e2"))
        pdf.setFillColor(HexColor("#f4f7fc"))
        pdf.roundRect(x, y, w, 34, 5, fill=1, stroke=1)
        pdf.setFillColor(HexColor("#17253d"))
        pdf.setFont("Helvetica", 9)
        pdf.drawCentredString(x + w / 2, y + 13, label)

    def arrow(x1, y1, x2, y2):
        pdf.setStrokeColor(HexColor("#647897"))
        pdf.line(x1, y1, x2, y2)
        angle = math.atan2(y2 - y1, x2 - x1)
        for delta in (-0.5, 0.5):
            pdf.line(x2, y2, x2 - 6 * math.cos(angle + delta), y2 - 6 * math.sin(angle + delta))

    for args in [(40, 135, 70, "Browser"), (135, 135, 70, "NGINX"), (230, 135, 80, "API replicas"),
                 (430, 135, 125, "Postgres / Valkey / S3"), (135, 65, 70, "Keycloak"),
                 (320, 65, 90, "Worker replicas"), (450, 65, 105, "Google egress")]:
        box(*args)
    for args in [(110, 152, 135, 152), (205, 152, 230, 152), (310, 152, 430, 152),
                 (170, 135, 170, 99), (365, 99, 453, 135), (410, 82, 450, 82)]:
        arrow(*args)
    pdf.showPage()
    title(2, "Security, recovery and qualification")
    block(
        705,
        "Security controls",
        [
            "Organization-scoped authorization plus PostgreSQL RLS under non-owner runtime roles.",
            "CSRF/origin checks, server-side sessions, role-gated reviews and Secure cookies in TLS mode.",
            "Bounded uploads/parsing, credential redaction, no live-device access or arbitrary agent tools.",
            "Shared limits, separate internal networks and verified TLS to fixed Google upstream hosts.",
            "Runtime roles append events but cannot rewrite history; database administrators remain trusted.",
        ],
    )
    block(
        569,
        "Fault tolerance and recovery",
        [
            "60-second configured job lease, 15-second heartbeat, fencing and bounded retries; at-least-once effects.",
            "Valkey outage refuses expensive work while allowing tightly limited evidence browsing.",
            "Consistent encrypted backup includes application/identity databases and encrypted objects.",
            "Restore into fresh volumes verifies record counts and every referenced artifact hash.",
            "One local host is one failure domain; no cloud deployment or uptime SLA is claimed.",
        ],
    )
    block(
        433,
        "Recorded qualification",
        [
            "Real browser tests cover OIDC, cross-tenant access, roles, mapping activation/rollback and PDFs.",
            f"Worker recovery: {measured('worker-recovery-result.json', 'recovery_seconds')} seconds; accepted work completed.",
            f"Restore rehearsal: {measured('restore-result.json', 'seconds')} seconds; plaintext artifact integrity checked.",
            f"Deterministic throughput: {measured('benchmark-result.json', 'configs_per_minute')} configurations/minute.",
            f"One-hour soak: {measured('soak-final.json', 'statuses.200')} successful reads; p95 {measured('soak-final.json', 'p95_ms')} ms. See docs/verification.md.",
        ],
    )
    block(
        287,
        "Remaining external validation",
        [
            "Independent expert-labelled vendor/firmware cases and licensed exact benchmark content.",
            "Controlled comparisons and ablations before any accuracy or novelty-superiority claim.",
            "Operational validation of proposed changes; no device commands are automatically applied.",
            "Future hosting requires real membership lifecycle, key rotation and host-level redundancy.",
            "Upstream image findings remain open; assessed exposure is in docs/dependency-review.md.",
        ],
    )
    pdf.save()


def presentation():
    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    content = [
        (
            "Evidence before verdicts",
            "Prooflane / Multi-vendor network security compliance auditing",
            [
                "Heterogeneous syntax creates consequential uncertainty.",
                "The application keeps unknown settings visible.",
                "The central contribution: tested, reviewed adaptation to unfamiliar formats.",
            ],
        ),
        (
            "One durable audit pipeline",
            "Local Docker services with independent API and worker replicas",
            [
                "NGINX + React -> FastAPI -> PostgreSQL, Valkey and encrypted S3",
                "Keycloak OIDC / PKCE -> tenant-scoped sessions and PostgreSQL RLS",
                "Leased worker jobs -> deterministic findings -> separate PDF stage",
                "Restricted Gemini proxy; optional metrics, traces and dashboards",
            ],
        ),
        (
            "Teach meaning; verify the change",
            "Agentic investigation with a deterministic approval boundary",
            [
                "Retrieve approved, versioned reference passages and identify material unknowns.",
                "Propose a constrained mapping; test positive and negative cases.",
                "A reviewer inspects semantics and labels before activation.",
                "Re-audit without redeployment; retire a mapping to roll back future use.",
            ],
        ),
        (
            "Qualification is observable",
            "Measured local results, with source evidence retained",
            [
                "Two-replica rate limit: 30 accepted / 25 rejected in the bounded burst test.",
                "Valkey outage: reads available, writes refused. API replica-loss reads succeeded.",
                f"Worker recovery: {measured('worker-recovery-result.json', 'recovery_seconds')} seconds; accepted work completed.",
                f"Fresh-volume restore: {measured('restore-result.json', 'seconds')} seconds; hashes verified.",
                f"{measured('benchmark-result.json', 'configs_per_minute')} configs/min; one-hour metadata p95 {measured('soak-final.json', 'p95_ms')} ms ({measured('soak-final.json', 'statuses.200')} successful reads).",
            ],
        ),
        (
            "An honest scope for a serious tool",
            "Working software, reviewable evidence and explicit limits",
            [
                "20 team technical checks; declared IOS, Junos set and FortiOS subsets.",
                "Synthetic examples; independently reviewed vendor labels remain future work.",
                "Candidate framework crosswalks; no certification or universal vendor claim.",
                "No automatic changes to live devices. One host is one failure domain.",
                "Open: upstream image findings; expert labels, exact editions and fair comparisons.",
            ],
        ),
    ]
    for number, (heading, subtitle, bullets) in enumerate(content, 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("F5F7FC")

        def text(x, y, w, h, value, size, color="17253D", bold=False):
            box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
            frame = box.text_frame
            frame.word_wrap = True
            paragraph = frame.paragraphs[0]
            paragraph.text = value
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(color)

        text(0.65, 0.35, 11, 0.4, "PROOFLANE  /  EVIDENCE CONSOLE", 12, "375C9D", True)
        text(0.65, 1.05, 12, 1.05, heading, 34, bold=True)
        text(0.65, 2.05, 12, 0.7, subtitle, 19, "52627A")
        for i, bullet in enumerate(bullets):
            text(0.85, 3.1 + i * 0.65, 11.8, 0.6, "• " + bullet, 18)
        text(
            0.65,
            7,
            11,
            0.3,
            "Synthetic-data qualification • Full conditions and limitations in docs/verification.md",
            10,
            "52627A",
        )
        text(12, 7, 0.6, 0.3, str(number) + " / 5", 10, "52627A")
    deck.save(OUT / "technical-presentation.pptx")


if __name__ == "__main__":
    architecture()
    presentation()
    print("Generated docs/deliverables/architecture.pdf (2 pages) and technical-presentation.pptx (5 slides)")
